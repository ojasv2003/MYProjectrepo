#!/usr/bin/env python3
"""Origin-destination pair-wise, month-wise yield analysis.

Yield is freight divided by charged weight:
  base yield  = basic freight / charged weight
  total yield = total freight / charged weight

Pair-month yields are volume-weighted: sum(freight) / sum(weight).
"""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent
DATA_PATH = ROOT / "data" / "Book_2_vas.xlsx"
OUTPUT_DIR = ROOT / "output"
EXCEL_PATH = OUTPUT_DIR / "OD_Yield_Analysis.xlsx"
HTML_PATH = OUTPUT_DIR / "od_yield_dashboard.html"
CSV_PATH = OUTPUT_DIR / "od_month_yield.csv"

MONTH_ORDER = ["April", "May", "June", "July"]
ZONE_ORDER = [
    "NORTH ZONE RO",
    "EAST ZONE RO",
    "WEST ZONE RO",
    "SOUTH ZONE RO",
    "NORTH EAST ZONE RO",
    "NEPAL ZONE",
]

NAVY = "1B365D"
TEAL = "0F6C8C"
WHITE = "FFFFFF"
GRID = "D7DEE8"
RED = "9B2C2C"
GREEN = "1F7A4D"

COL_LABELS = {
    "origin_zone": "Origin zone",
    "destination_zone": "Destination zone",
    "od_pair": "OD pair",
    "month": "Month",
    "shipments": "Shipments",
    "charged_weight": "Charged weight",
    "basic_freight": "Basic freight",
    "total_freight": "Total freight",
    "vas_freight": "VAS freight",
    "base_yield": "Base yield",
    "total_yield": "Total yield",
    "vas_yield": "VAS yield",
    "yield_uplift": "Yield uplift (total − base)",
    "avg_shipment_base_yield": "Avg shipment base yield",
    "avg_shipment_total_yield": "Avg shipment total yield",
    "avg_weight_per_shipment": "Avg weight / shipment",
    "weight_share_pct": "Weight share %",
    "basic_freight_share_pct": "Basic freight share %",
    "total_freight_share_pct": "Total freight share %",
    "base_yield_mom_abs": "Base yield MoM abs",
    "base_yield_mom_pct": "Base yield MoM %",
    "total_yield_mom_abs": "Total yield MoM abs",
    "total_yield_mom_pct": "Total yield MoM %",
    "charged_weight_mom_pct": "Weight MoM %",
    "weight_mom_pct": "Weight MoM %",
    "weight_rank": "Weight rank",
    "base_yield_vs_network": "Base yield vs network",
    "total_yield_vs_network": "Total yield vs network",
    "All_months": "All months",
    "issue": "Issue",
    "base_yield_shipment": "Shipment base yield",
    "total_yield_shipment": "Shipment total yield",
}


def labeled(df: pd.DataFrame) -> pd.DataFrame:
    return df.rename(columns={k: v for k, v in COL_LABELS.items() if k in df.columns})


def labeled_fmt(fmt: dict[str, str]) -> dict[str, str]:
    return {COL_LABELS.get(k, k): v for k, v in fmt.items()}


def load_shipments(path: Path) -> pd.DataFrame:
    df = pd.read_excel(path, header=1)
    df = df.loc[:, ~df.columns.astype(str).str.startswith("Unnamed")]
    df.columns = [
        "month",
        "origin_zone",
        "destination_zone",
        "charged_weight",
        "basic_freight",
        "total_freight",
    ]
    df["month"] = pd.Categorical(df["month"], categories=MONTH_ORDER, ordered=True)
    df["origin_zone"] = pd.Categorical(df["origin_zone"], categories=ZONE_ORDER, ordered=True)
    df["destination_zone"] = pd.Categorical(
        df["destination_zone"], categories=ZONE_ORDER, ordered=True
    )
    df["vas_freight"] = df["total_freight"] - df["basic_freight"]
    df["base_yield_shipment"] = df["basic_freight"] / df["charged_weight"]
    df["total_yield_shipment"] = df["total_freight"] / df["charged_weight"]
    return df


def summarize(df: pd.DataFrame, group_cols: list[str]) -> pd.DataFrame:
    g = (
        df.groupby(group_cols, observed=True, dropna=False)
        .agg(
            shipments=("charged_weight", "count"),
            charged_weight=("charged_weight", "sum"),
            basic_freight=("basic_freight", "sum"),
            total_freight=("total_freight", "sum"),
            vas_freight=("vas_freight", "sum"),
            avg_shipment_base_yield=("base_yield_shipment", "mean"),
            avg_shipment_total_yield=("total_yield_shipment", "mean"),
        )
        .reset_index()
    )
    g["base_yield"] = g["basic_freight"] / g["charged_weight"]
    g["total_yield"] = g["total_freight"] / g["charged_weight"]
    g["vas_yield"] = g["vas_freight"] / g["charged_weight"]
    g["yield_uplift"] = g["total_yield"] - g["base_yield"]
    g["avg_weight_per_shipment"] = g["charged_weight"] / g["shipments"]
    return g


def add_shares(g: pd.DataFrame, total_weight: float, total_basic: float, total_freight: float) -> pd.DataFrame:
    g = g.copy()
    g["weight_share_pct"] = 100.0 * g["charged_weight"] / total_weight
    g["basic_freight_share_pct"] = 100.0 * g["basic_freight"] / total_basic
    g["total_freight_share_pct"] = 100.0 * g["total_freight"] / total_freight
    return g


def add_mom(od_month: pd.DataFrame) -> pd.DataFrame:
    out = od_month.sort_values(["origin_zone", "destination_zone", "month"]).copy()
    keys = ["origin_zone", "destination_zone"]
    for col in ["base_yield", "total_yield", "vas_yield", "charged_weight", "shipments"]:
        prev = out.groupby(keys, observed=True)[col].shift(1)
        out[f"{col}_mom_abs"] = out[col] - prev
        out[f"{col}_mom_pct"] = 100.0 * (out[col] - prev) / prev.replace(0, pd.NA)
    return out


def pivot_metric(od_month: pd.DataFrame, value: str) -> pd.DataFrame:
    p = od_month.pivot_table(
        index=["origin_zone", "destination_zone"],
        columns="month",
        values=value,
        observed=True,
        aggfunc="first",
    )
    p = p.reindex(columns=MONTH_ORDER)
    p["All months"] = None
    return p.reset_index()


def style_header(ws, row: int, cols: int) -> None:
    fill = PatternFill("solid", fgColor=NAVY)
    font = Font(name="Calibri", bold=True, color=WHITE, size=11)
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = align
    ws.row_dimensions[row].height = 32


def apply_border(ws, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    thin = Border(
        left=Side(style="thin", color=GRID),
        right=Side(style="thin", color=GRID),
        top=Side(style="thin", color=GRID),
        bottom=Side(style="thin", color=GRID),
    )
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            ws.cell(row=r, column=c).border = thin
            ws.cell(row=r, column=c).font = Font(name="Calibri", size=10)


def autosize(ws, min_width: int = 12, max_width: int = 28) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = 0
        for cell in col[:80]:
            if cell.value is None:
                continue
            length = max(length, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(length + 2, min_width), max_width)


def write_df(
    ws,
    df: pd.DataFrame,
    start_row: int = 1,
    number_formats: dict[str, str] | None = None,
    table_name: str | None = None,
    color_scale_cols: list[str] | None = None,
) -> None:
    df = labeled(df)
    number_formats = labeled_fmt(number_formats or {})
    color_scale_cols = [COL_LABELS.get(c, c) for c in (color_scale_cols or [])]
    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), start=start_row):
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=None if pd.isna(value) else value)
            if r_idx == start_row:
                continue
            col_name = df.columns[c_idx - 1]
            if col_name in number_formats and isinstance(value, (int, float)) and pd.notna(value):
                cell.number_format = number_formats[col_name]
    header_row = start_row
    data_end = start_row + len(df)
    cols = len(df.columns)
    style_header(ws, header_row, cols)
    apply_border(ws, header_row, data_end, 1, cols)
    ws.freeze_panes = f"A{header_row + 1}"
    if table_name and len(df) > 0:
        ref = f"A{header_row}:{get_column_letter(cols)}{data_end}"
        table = Table(displayName=table_name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    else:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(cols)}{data_end}"
    if color_scale_cols:
        for name in color_scale_cols:
            if name not in df.columns:
                continue
            idx = list(df.columns).index(name) + 1
            letter = get_column_letter(idx)
            ws.conditional_formatting.add(
                f"{letter}{header_row + 1}:{letter}{data_end}",
                ColorScaleRule(
                    start_type="min",
                    start_color="F4C7C3",
                    mid_type="percentile",
                    mid_value=50,
                    mid_color="FFF2CC",
                    end_type="max",
                    end_color="B7E1CD",
                ),
            )
    autosize(ws)


def title_block(ws, title: str, subtitle: str, cols: int = 8) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=cols)
    c1 = ws.cell(row=1, column=1, value=title)
    c1.font = Font(name="Calibri", bold=True, size=16, color=NAVY)
    c2 = ws.cell(row=2, column=1, value=subtitle)
    c2.font = Font(name="Calibri", size=11, italic=True, color=TEAL)
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 18


def build_excel(
    shipments: pd.DataFrame,
    od_month: pd.DataFrame,
    od_all: pd.DataFrame,
    month_sum: pd.DataFrame,
    origin_month: pd.DataFrame,
    dest_month: pd.DataFrame,
    quality: pd.DataFrame,
) -> None:
    wb = Workbook()
    fmt_int = "#,##0"
    fmt_qty = "#,##0.00"
    fmt_y = "0.0000"
    fmt_pct = "0.00"
    fmt_mom = "0.00"

    # --- Notes ---
    notes = wb.active
    notes.title = "Notes"
    notes["A1"] = "Origin–Destination pair-wise, month-wise yield analysis"
    notes["A1"].font = Font(name="Calibri", bold=True, size=18, color=NAVY)
    notes.merge_cells("A1:B1")
    bullets = [
        ("Source file", "Book_2_vas.xlsx (Sheet1)"),
        ("Grain", "One row per shipment; analysis rolled up to Origin × Destination × Month"),
        ("Months", "April, May, June, July"),
        ("Base yield", "Basic freight ÷ Charged weight"),
        ("Total yield", "Total freight ÷ Charged weight"),
        ("VAS yield", "(Total freight − Basic freight) ÷ Charged weight"),
        ("Pair-month yield", "Volume-weighted: Σ freight / Σ weight (not the average of shipment yields)"),
        ("Avg shipment yield", "Unweighted mean of shipment-level freight/weight, shown for comparison only"),
        ("Yield uplift", "Total yield − Base yield (VAS contribution per unit weight)"),
        ("MoM change", "Versus the previous month for the same origin–destination pair"),
        ("Zero freight", "Shipments with zero basic/total freight are retained; they lower realized yield"),
        ("Sparse pairs", "Some OD pairs have few shipments; interpret their yield with the shipment count"),
    ]
    notes["A3"] = "Field"
    notes["B3"] = "Definition"
    style_header(notes, 3, 2)
    for i, (k, v) in enumerate(bullets, start=4):
        notes.cell(row=i, column=1, value=k).font = Font(name="Calibri", bold=True)
        notes.cell(row=i, column=2, value=v).font = Font(name="Calibri")
        notes.cell(row=i, column=2).alignment = Alignment(wrap_text=True)
    notes.column_dimensions["A"].width = 24
    notes.column_dimensions["B"].width = 110
    notes.row_dimensions[1].height = 26

    overall_w = shipments["charged_weight"].sum()
    overall_b = shipments["basic_freight"].sum()
    overall_t = shipments["total_freight"].sum()
    kpis = [
        ("Shipments", f"{len(shipments):,}"),
        ("OD pairs", f"{shipments.groupby(['origin_zone','destination_zone'], observed=True).ngroups}"),
        ("OD × month combinations", f"{len(od_month)}"),
        ("Charged weight", f"{overall_w:,.1f}"),
        ("Basic freight", f"{overall_b:,.2f}"),
        ("Total freight", f"{overall_t:,.2f}"),
        ("Network base yield", f"{overall_b / overall_w:.4f}"),
        ("Network total yield", f"{overall_t / overall_w:.4f}"),
        ("Network VAS yield", f"{(overall_t - overall_b) / overall_w:.4f}"),
    ]
    notes["A17"] = "Network snapshot"
    notes["A17"].font = Font(name="Calibri", bold=True, size=14, color=NAVY)
    notes["A18"] = "Metric"
    notes["B18"] = "Value"
    style_header(notes, 18, 2)
    for i, (k, v) in enumerate(kpis, start=19):
        notes.cell(row=i, column=1, value=k)
        notes.cell(row=i, column=2, value=v)

    # --- Insights ---
    insights = wb.create_sheet("Insights")
    title_block(
        insights,
        "Key findings — OD pair × month yield",
        "Yield = freight ÷ charged weight. Pair-month figures are volume-weighted (Σ freight / Σ weight).",
        cols=6,
    )
    net_base = overall_b / overall_w
    net_total = overall_t / overall_w
    month_sorted = month_sum.sort_values("month")
    peak_total = month_sorted.loc[month_sorted["total_yield"].idxmax()]
    trough_total = month_sorted.loc[month_sorted["total_yield"].idxmin()]
    peak_base = month_sorted.loc[month_sorted["base_yield"].idxmax()]
    trough_base = month_sorted.loc[month_sorted["base_yield"].idxmin()]
    od_ranked = od_all.copy()
    od_ranked["od_pair"] = (
        od_ranked["origin_zone"].astype(str) + " → " + od_ranked["destination_zone"].astype(str)
    )
    od_ranked = od_ranked.sort_values("charged_weight", ascending=False)
    top5_share = od_ranked.head(5)["charged_weight"].sum() / overall_w * 100
    largest = od_ranked.iloc[0]
    large = od_ranked[od_ranked["weight_share_pct"] >= 1.0]
    high_total = large.sort_values("total_yield", ascending=False).iloc[0]
    low_total = large.sort_values("total_yield", ascending=True).iloc[0]
    high_vas = large.sort_values("vas_yield", ascending=False).iloc[0]
    below_net = od_ranked[od_ranked["base_yield"] < net_base].head(3)

    june = month_sorted[month_sorted["month"] == "June"].iloc[0]
    july = month_sorted[month_sorted["month"] == "July"].iloc[0]
    july_base_drop = 100.0 * (july["base_yield"] - june["base_yield"]) / june["base_yield"]
    july_total_drop = 100.0 * (july["total_yield"] - june["total_yield"]) / june["total_yield"]

    findings = [
        (
            "Network yield",
            f"Across {len(shipments):,} shipments the network earned base yield {net_base:.4f} "
            f"and total yield {net_total:.4f} (VAS yield {(net_total - net_base):.4f}). "
            f"Total freight is {((overall_t / overall_b) - 1) * 100:.1f}% above basic freight.",
        ),
        (
            "Month trend",
            f"Total yield peaked in {peak_total['month']} at {peak_total['total_yield']:.4f} "
            f"and was lowest in {trough_total['month']} at {trough_total['total_yield']:.4f}. "
            f"Base yield peaked in {peak_base['month']} ({peak_base['base_yield']:.4f}) and "
            f"was lowest in {trough_base['month']} ({trough_base['base_yield']:.4f}). "
            f"July vs June: base yield {july_base_drop:+.1f}%, total yield {july_total_drop:+.1f}%.",
        ),
        (
            "Volume concentration",
            f"The top 5 OD pairs carry {top5_share:.1f}% of charged weight. "
            f"Largest pair is {largest['od_pair']} "
            f"({largest['weight_share_pct']:.1f}% of weight, {int(largest['shipments']):,} shipments) "
            f"with base yield {largest['base_yield']:.4f} and total yield {largest['total_yield']:.4f} "
            f"— {'below' if largest['base_yield'] < net_base else 'above'} the network base yield.",
        ),
        (
            "Highest-yield large pair",
            f"Among pairs with ≥1% of network weight, {high_total['od_pair']} has the highest total yield "
            f"({high_total['total_yield']:.4f}; base {high_total['base_yield']:.4f}).",
        ),
        (
            "Lowest-yield large pair",
            f"Among pairs with ≥1% of network weight, {low_total['od_pair']} has the lowest total yield "
            f"({low_total['total_yield']:.4f}; base {low_total['base_yield']:.4f}).",
        ),
        (
            "VAS / yield uplift",
            f"{high_vas['od_pair']} has the strongest VAS yield among large pairs "
            f"({high_vas['vas_yield']:.4f} extra freight per unit weight).",
        ),
        (
            "How to read the workbook",
            "OD_Month_Yield is the primary grain (origin × destination × month). "
            "Pivot sheets show the same yields in matrix form. Use Rankings only for pairs with "
            "at least 30 shipments that month. Data_Quality lists zero or inverted freight rows "
            "that remain in the totals and pull yield down.",
        ),
    ]
    insights["A4"] = "Theme"
    insights["B4"] = "Finding"
    style_header(insights, 4, 2)
    for i, (theme, text) in enumerate(findings, start=5):
        insights.cell(row=i, column=1, value=theme).font = Font(name="Calibri", bold=True, size=11)
        cell = insights.cell(row=i, column=2, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        insights.row_dimensions[i].height = 48
    insights.column_dimensions["A"].width = 28
    insights.column_dimensions["B"].width = 120

    insights["A13"] = "Largest OD pairs vs network yield"
    insights["A13"].font = Font(name="Calibri", bold=True, size=13, color=NAVY)
    top_compare = od_ranked.head(10)[
        [
            "od_pair",
            "shipments",
            "charged_weight",
            "weight_share_pct",
            "base_yield",
            "total_yield",
            "vas_yield",
        ]
    ].copy()
    top_compare["base_yield_vs_network"] = top_compare["base_yield"] - net_base
    top_compare["total_yield_vs_network"] = top_compare["total_yield"] - net_total
    write_df(
        insights,
        top_compare,
        start_row=14,
        number_formats={
            "shipments": fmt_int,
            "charged_weight": fmt_qty,
            "weight_share_pct": fmt_pct,
            "base_yield": fmt_y,
            "total_yield": fmt_y,
            "vas_yield": fmt_y,
            "base_yield_vs_network": fmt_y,
            "total_yield_vs_network": fmt_y,
        },
        table_name="tblTopPairsVsNetwork",
        color_scale_cols=["base_yield", "total_yield"],
    )
    insights.cell(row=26, column=1, value="High-volume pairs still below network base yield").font = Font(
        name="Calibri", bold=True, size=13, color=NAVY
    )
    write_df(
        insights,
        below_net[["od_pair", "shipments", "charged_weight", "weight_share_pct", "base_yield", "total_yield"]],
        start_row=27,
        number_formats={
            "shipments": fmt_int,
            "charged_weight": fmt_qty,
            "weight_share_pct": fmt_pct,
            "base_yield": fmt_y,
            "total_yield": fmt_y,
        },
        table_name="tblBelowNetwork",
    )

    # --- Main OD-month table ---
    main_cols = [
        "origin_zone",
        "destination_zone",
        "od_pair",
        "month",
        "shipments",
        "charged_weight",
        "basic_freight",
        "total_freight",
        "vas_freight",
        "base_yield",
        "total_yield",
        "vas_yield",
        "yield_uplift",
        "avg_shipment_base_yield",
        "avg_shipment_total_yield",
        "avg_weight_per_shipment",
        "weight_share_pct",
        "total_freight_share_pct",
        "base_yield_mom_abs",
        "base_yield_mom_pct",
        "total_yield_mom_abs",
        "total_yield_mom_pct",
        "charged_weight_mom_pct",
    ]
    od_view = od_month.copy()
    od_view["od_pair"] = od_view["origin_zone"].astype(str) + " → " + od_view["destination_zone"].astype(str)
    od_view = od_view[main_cols]
    od_view = od_view.sort_values(["origin_zone", "destination_zone", "month"])

    num_main = {
        "shipments": fmt_int,
        "charged_weight": fmt_qty,
        "basic_freight": fmt_qty,
        "total_freight": fmt_qty,
        "vas_freight": fmt_qty,
        "base_yield": fmt_y,
        "total_yield": fmt_y,
        "vas_yield": fmt_y,
        "yield_uplift": fmt_y,
        "avg_shipment_base_yield": fmt_y,
        "avg_shipment_total_yield": fmt_y,
        "avg_weight_per_shipment": fmt_qty,
        "weight_share_pct": fmt_pct,
        "total_freight_share_pct": fmt_pct,
        "base_yield_mom_abs": "0.0000",
        "base_yield_mom_pct": fmt_mom,
        "total_yield_mom_abs": "0.0000",
        "total_yield_mom_pct": fmt_mom,
        "charged_weight_mom_pct": fmt_mom,
    }

    ws_main = wb.create_sheet("OD_Month_Yield")
    title_block(
        ws_main,
        "Origin–Destination × Month yield",
        "Base yield = Σ basic freight / Σ charged weight; Total yield = Σ total freight / Σ charged weight",
        cols=len(main_cols),
    )
    write_df(
        ws_main,
        od_view,
        start_row=4,
        number_formats=num_main,
        table_name="tblODMonthYield",
        color_scale_cols=["base_yield", "total_yield", "vas_yield"],
    )

    # --- Pivots ---
    def write_pivot(name: str, title: str, metric: str, overall_map: pd.Series) -> None:
        p = od_month.pivot_table(
            index=["origin_zone", "destination_zone"],
            columns="month",
            values=metric,
            observed=True,
            aggfunc="first",
        ).reindex(columns=MONTH_ORDER)
        p = p.reset_index()
        p["od_pair"] = p["origin_zone"].astype(str) + " → " + p["destination_zone"].astype(str)
        p["All_months"] = p.apply(
            lambda r: overall_map.get((r["origin_zone"], r["destination_zone"])), axis=1
        )
        p = p[
            ["origin_zone", "destination_zone", "od_pair", "April", "May", "June", "July", "All_months"]
        ]
        ws = wb.create_sheet(name)
        title_block(ws, title, "Blank cells mean the OD pair had no shipments that month", cols=8)
        formats = {m: fmt_y if metric.endswith("yield") else (fmt_int if metric == "shipments" else fmt_qty) for m in MONTH_ORDER + ["All_months"]}
        write_df(
            ws,
            p,
            start_row=4,
            number_formats=formats,
            table_name="tbl" + name.replace("_", ""),
            color_scale_cols=MONTH_ORDER + ["All_months"],
        )

    od_all_idx = od_all.set_index(["origin_zone", "destination_zone"])
    write_pivot("Pivot_Base_Yield", "Base yield by OD pair and month", "base_yield", od_all_idx["base_yield"])
    write_pivot("Pivot_Total_Yield", "Total yield by OD pair and month", "total_yield", od_all_idx["total_yield"])
    write_pivot("Pivot_VAS_Yield", "VAS yield by OD pair and month", "vas_yield", od_all_idx["vas_yield"])
    write_pivot("Pivot_Weight", "Charged weight by OD pair and month", "charged_weight", od_all_idx["charged_weight"])

    # --- OD overall ---
    od_all_view = od_all.copy()
    od_all_view["od_pair"] = (
        od_all_view["origin_zone"].astype(str) + " → " + od_all_view["destination_zone"].astype(str)
    )
    od_all_view["base_yield_vs_network"] = od_all_view["base_yield"] - (overall_b / overall_w)
    od_all_view["total_yield_vs_network"] = od_all_view["total_yield"] - (overall_t / overall_w)
    od_all_view = od_all_view.sort_values("charged_weight", ascending=False)
    od_all_view["weight_rank"] = range(1, len(od_all_view) + 1)
    od_all_cols = [
        "weight_rank",
        "od_pair",
        "origin_zone",
        "destination_zone",
        "shipments",
        "charged_weight",
        "basic_freight",
        "total_freight",
        "vas_freight",
        "base_yield",
        "total_yield",
        "vas_yield",
        "yield_uplift",
        "weight_share_pct",
        "total_freight_share_pct",
        "base_yield_vs_network",
        "total_yield_vs_network",
        "avg_weight_per_shipment",
    ]
    ws_od = wb.create_sheet("OD_Overall")
    title_block(ws_od, "OD pair yield across all months", "Ranked by charged weight", cols=len(od_all_cols))
    write_df(
        ws_od,
        od_all_view[od_all_cols],
        start_row=4,
        number_formats={
            "weight_rank": "0",
            "shipments": fmt_int,
            "charged_weight": fmt_qty,
            "basic_freight": fmt_qty,
            "total_freight": fmt_qty,
            "vas_freight": fmt_qty,
            "base_yield": fmt_y,
            "total_yield": fmt_y,
            "vas_yield": fmt_y,
            "yield_uplift": fmt_y,
            "weight_share_pct": fmt_pct,
            "total_freight_share_pct": fmt_pct,
            "base_yield_vs_network": "0.0000",
            "total_yield_vs_network": "0.0000",
            "avg_weight_per_shipment": fmt_qty,
        },
        table_name="tblODOverall",
        color_scale_cols=["base_yield", "total_yield"],
    )

    # --- Month summary ---
    month_view = month_sum.copy()
    month_view["base_yield_mom_pct"] = 100.0 * month_view["base_yield"].pct_change()
    month_view["total_yield_mom_pct"] = 100.0 * month_view["total_yield"].pct_change()
    month_view["weight_mom_pct"] = 100.0 * month_view["charged_weight"].pct_change()
    ws_m = wb.create_sheet("Month_Summary")
    title_block(ws_m, "Network yield by month", "All origin–destination pairs combined", cols=12)
    write_df(
        ws_m,
        month_view,
        start_row=4,
        number_formats={
            "shipments": fmt_int,
            "charged_weight": fmt_qty,
            "basic_freight": fmt_qty,
            "total_freight": fmt_qty,
            "vas_freight": fmt_qty,
            "base_yield": fmt_y,
            "total_yield": fmt_y,
            "vas_yield": fmt_y,
            "yield_uplift": fmt_y,
            "avg_shipment_base_yield": fmt_y,
            "avg_shipment_total_yield": fmt_y,
            "avg_weight_per_shipment": fmt_qty,
            "weight_share_pct": fmt_pct,
            "basic_freight_share_pct": fmt_pct,
            "total_freight_share_pct": fmt_pct,
            "base_yield_mom_pct": fmt_mom,
            "total_yield_mom_pct": fmt_mom,
            "weight_mom_pct": fmt_mom,
        },
        table_name="tblMonthSummary",
    )

    line = LineChart()
    line.title = "Network yield by month"
    line.y_axis.title = "Yield (freight / weight)"
    line.x_axis.title = "Month"
    line.height = 8
    line.width = 15
    line.style = 10
    labeled_month_cols = [COL_LABELS.get(c, c) for c in month_view.columns]
    col_index = {c: i + 1 for i, c in enumerate(labeled_month_cols)}
    cats = Reference(ws_m, min_col=col_index["Month"], min_row=5, max_row=4 + len(month_view))
    y1 = Reference(
        ws_m,
        min_col=col_index["Base yield"],
        min_row=4,
        max_row=4 + len(month_view),
    )
    y2 = Reference(
        ws_m,
        min_col=col_index["Total yield"],
        min_row=4,
        max_row=4 + len(month_view),
    )
    line.add_data(y1, titles_from_data=True)
    line.add_data(y2, titles_from_data=True)
    line.set_categories(cats)
    line.shape = 4
    ws_m.add_chart(line, "A12")

    # --- Origin / Dest month ---
    for sheet_name, frame, title in [
        ("Origin_Month", origin_month, "Yield by origin zone and month"),
        ("Destination_Month", dest_month, "Yield by destination zone and month"),
    ]:
        ws = wb.create_sheet(sheet_name)
        title_block(ws, title, "Volume-weighted yields", cols=len(frame.columns))
        write_df(
            ws,
            frame,
            start_row=4,
            number_formats={
                "shipments": fmt_int,
                "charged_weight": fmt_qty,
                "basic_freight": fmt_qty,
                "total_freight": fmt_qty,
                "vas_freight": fmt_qty,
                "base_yield": fmt_y,
                "total_yield": fmt_y,
                "vas_yield": fmt_y,
                "yield_uplift": fmt_y,
                "avg_shipment_base_yield": fmt_y,
                "avg_shipment_total_yield": fmt_y,
                "avg_weight_per_shipment": fmt_qty,
                "weight_share_pct": fmt_pct,
                "basic_freight_share_pct": fmt_pct,
                "total_freight_share_pct": fmt_pct,
            },
        table_name="tbl" + sheet_name.replace("_", ""),
            color_scale_cols=["base_yield", "total_yield"],
        )

    # --- Rankings ---
    min_shipments = 30
    ranked = od_month[od_month["shipments"] >= min_shipments].copy()
    ranked["od_pair"] = ranked["origin_zone"].astype(str) + " → " + ranked["destination_zone"].astype(str)
    top_base = ranked.nlargest(15, "base_yield")[
        ["month", "od_pair", "shipments", "charged_weight", "base_yield", "total_yield"]
    ]
    bot_base = ranked.nsmallest(15, "base_yield")[
        ["month", "od_pair", "shipments", "charged_weight", "base_yield", "total_yield"]
    ]
    top_total = ranked.nlargest(15, "total_yield")[
        ["month", "od_pair", "shipments", "charged_weight", "base_yield", "total_yield"]
    ]
    bot_total = ranked.nsmallest(15, "total_yield")[
        ["month", "od_pair", "shipments", "charged_weight", "base_yield", "total_yield"]
    ]
    ws_r = wb.create_sheet("Rankings")
    ws_r["A1"] = "Yield rankings (OD × month with at least 30 shipments)"
    ws_r["A1"].font = Font(name="Calibri", bold=True, size=16, color=NAVY)
    ws_r["A3"] = "Highest base yield"
    ws_r["A3"].font = Font(name="Calibri", bold=True, size=12, color=GREEN)
    write_df(ws_r, top_base, start_row=4, number_formats={"shipments": fmt_int, "charged_weight": fmt_qty, "base_yield": fmt_y, "total_yield": fmt_y})
    start_bot = 4 + len(top_base) + 3
    ws_r.cell(row=start_bot, column=1, value="Lowest base yield").font = Font(
        name="Calibri", bold=True, size=12, color=RED
    )
    write_df(
        ws_r,
        bot_base,
        start_row=start_bot + 1,
        number_formats={"shipments": fmt_int, "charged_weight": fmt_qty, "base_yield": fmt_y, "total_yield": fmt_y},
    )
    ws_r["H3"] = "Highest total yield"
    ws_r["H3"].font = Font(name="Calibri", bold=True, size=12, color=GREEN)
    # write second block starting at column H by manual copy
    start_col = 8
    raw_headers = list(top_total.columns)
    headers = [COL_LABELS.get(c, c) for c in raw_headers]
    fill = PatternFill("solid", fgColor=NAVY)
    font = Font(name="Calibri", bold=True, color=WHITE, size=11)
    for c, h in enumerate(headers, start=start_col):
        cell = ws_r.cell(row=4, column=c, value=h)
        cell.fill = fill
        cell.font = font
    for r, rec in enumerate(top_total.itertuples(index=False), start=5):
        for c, val in enumerate(rec, start=start_col):
            cell = ws_r.cell(row=r, column=c, value=val)
            name = raw_headers[c - start_col]
            if name in ("base_yield", "total_yield"):
                cell.number_format = fmt_y
            elif name == "shipments":
                cell.number_format = fmt_int
            elif name == "charged_weight":
                cell.number_format = fmt_qty

    ws_r.cell(row=start_bot, column=8, value="Lowest total yield").font = Font(
        name="Calibri", bold=True, size=12, color=RED
    )
    for c, h in enumerate(headers, start=start_col):
        cell = ws_r.cell(row=start_bot + 1, column=c, value=h)
        cell.fill = fill
        cell.font = font
    for r, rec in enumerate(bot_total.itertuples(index=False), start=start_bot + 2):
        for c, val in enumerate(rec, start=start_col):
            cell = ws_r.cell(row=r, column=c, value=val)
            name = raw_headers[c - start_col]
            if name in ("base_yield", "total_yield"):
                cell.number_format = fmt_y
            elif name == "shipments":
                cell.number_format = fmt_int
            elif name == "charged_weight":
                cell.number_format = fmt_qty

    bar = BarChart()
    bar.type = "col"
    bar.title = "Top 10 OD pairs by charged weight — base vs total yield"
    bar.y_axis.title = "Yield"
    top10 = od_all_view.head(10)[["od_pair", "base_yield", "total_yield"]]
    helper_row = 42
    ws_r.cell(row=helper_row, column=15, value="OD pair")
    ws_r.cell(row=helper_row, column=16, value="Base yield")
    ws_r.cell(row=helper_row, column=17, value="Total yield")
    for i, rec in enumerate(top10.itertuples(index=False), start=helper_row + 1):
        ws_r.cell(row=i, column=15, value=rec.od_pair)
        ws_r.cell(row=i, column=16, value=rec.base_yield)
        ws_r.cell(row=i, column=17, value=rec.total_yield)
    data = Reference(ws_r, min_col=16, min_row=helper_row, max_col=17, max_row=helper_row + len(top10))
    cats = Reference(ws_r, min_col=15, min_row=helper_row + 1, max_row=helper_row + len(top10))
    bar.add_data(data, titles_from_data=True)
    bar.set_categories(cats)
    bar.shape = 4
    bar.height = 10
    bar.width = 18
    ws_r.add_chart(bar, "A28")

    # --- Data quality ---
    ws_q = wb.create_sheet("Data_Quality")
    title_block(
        ws_q,
        "Shipments with zero or inverted freight",
        "Included in yield totals; listed here so they can be reviewed",
        cols=8,
    )
    write_df(ws_q, quality, start_row=4, table_name="tblDataQuality")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(EXCEL_PATH)


def build_html(od_month: pd.DataFrame, od_all: pd.DataFrame, month_sum: pd.DataFrame, shipments: pd.DataFrame) -> None:
    def recs(df: pd.DataFrame) -> list[dict]:
        out = []
        for row in df.itertuples(index=False):
            item = {}
            for col in df.columns:
                val = getattr(row, col)
                if pd.isna(val):
                    item[col] = None
                elif hasattr(val, "item"):
                    item[col] = val.item()
                else:
                    item[col] = val if not isinstance(val, (pd.Timestamp,)) else str(val)
            out.append(item)
        return out

    payload = {
        "months": MONTH_ORDER,
        "zones": ZONE_ORDER,
        "network": {
            "shipments": int(len(shipments)),
            "weight": float(shipments["charged_weight"].sum()),
            "basic": float(shipments["basic_freight"].sum()),
            "total": float(shipments["total_freight"].sum()),
        },
        "od_month": recs(
            od_month.assign(
                od_pair=od_month["origin_zone"].astype(str)
                + " → "
                + od_month["destination_zone"].astype(str)
            )[
                [
                    "origin_zone",
                    "destination_zone",
                    "od_pair",
                    "month",
                    "shipments",
                    "charged_weight",
                    "basic_freight",
                    "total_freight",
                    "vas_freight",
                    "base_yield",
                    "total_yield",
                    "vas_yield",
                    "yield_uplift",
                    "weight_share_pct",
                    "base_yield_mom_pct",
                    "total_yield_mom_pct",
                ]
            ]
        ),
        "od_all": recs(
            od_all.assign(
                od_pair=od_all["origin_zone"].astype(str)
                + " → "
                + od_all["destination_zone"].astype(str)
            )[
                [
                    "origin_zone",
                    "destination_zone",
                    "od_pair",
                    "shipments",
                    "charged_weight",
                    "basic_freight",
                    "total_freight",
                    "base_yield",
                    "total_yield",
                    "vas_yield",
                    "weight_share_pct",
                ]
            ]
        ),
        "months_summary": recs(
            month_sum[
                [
                    "month",
                    "shipments",
                    "charged_weight",
                    "basic_freight",
                    "total_freight",
                    "base_yield",
                    "total_yield",
                    "vas_yield",
                ]
            ]
        ),
    }

    html = r"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8" />
<meta name="viewport" content="width=device-width, initial-scale=1" />
<title>OD Pair Month-wise Yield Analysis</title>
<style>
  :root {
    --navy: #1b365d;
    --teal: #0f6c8c;
    --gold: #c4a35a;
    --bg: #eef3f8;
    --card: #ffffff;
    --ink: #1a2332;
    --muted: #5b6b7c;
    --line: #d5dee8;
    --good: #1f7a4d;
    --bad: #9b2c2c;
  }
  * { box-sizing: border-box; }
  body {
    margin: 0; font-family: "Segoe UI", Calibri, sans-serif;
    background: var(--bg); color: var(--ink);
  }
  header {
    background: linear-gradient(120deg, var(--navy), var(--teal));
    color: #fff; padding: 28px 32px 22px;
  }
  header h1 { margin: 0 0 6px; font-size: 26px; font-weight: 650; }
  header p { margin: 0; opacity: .9; max-width: 900px; }
  .wrap { padding: 20px 24px 48px; max-width: 1400px; margin: 0 auto; }
  .kpis { display: grid; grid-template-columns: repeat(auto-fit, minmax(180px, 1fr)); gap: 12px; margin: -18px 0 18px; }
  .kpi {
    background: var(--card); border-radius: 12px; padding: 14px 16px;
    box-shadow: 0 6px 18px rgba(27,54,93,.08); border: 1px solid var(--line);
  }
  .kpi .lbl { font-size: 12px; color: var(--muted); text-transform: uppercase; letter-spacing: .04em; }
  .kpi .val { font-size: 22px; font-weight: 700; color: var(--navy); margin-top: 4px; }
  .kpi .sub { font-size: 12px; color: var(--muted); margin-top: 2px; }
  .panel {
    background: var(--card); border-radius: 12px; padding: 16px 18px 18px;
    box-shadow: 0 6px 18px rgba(27,54,93,.08); border: 1px solid var(--line); margin-bottom: 16px;
  }
  .panel h2 { margin: 0 0 12px; font-size: 16px; color: var(--navy); }
  .filters { display: flex; flex-wrap: wrap; gap: 10px; align-items: end; }
  label { display: flex; flex-direction: column; font-size: 12px; color: var(--muted); gap: 4px; }
  select, input {
    min-width: 180px; padding: 8px 10px; border: 1px solid var(--line);
    border-radius: 8px; font-size: 14px; background: #fff;
  }
  button {
    background: var(--navy); color: #fff; border: 0; border-radius: 8px;
    padding: 9px 14px; cursor: pointer; font-size: 14px;
  }
  button.secondary { background: #fff; color: var(--navy); border: 1px solid var(--line); }
  .heat { overflow: auto; }
  table { border-collapse: collapse; width: 100%; font-size: 13px; }
  th, td { border-bottom: 1px solid var(--line); padding: 7px 8px; text-align: right; white-space: nowrap; }
  th { background: var(--navy); color: #fff; position: sticky; top: 0; font-weight: 600; cursor: pointer; }
  td.l, th.l { text-align: left; }
  tbody tr:hover { background: #f3f8fc; }
  .heat td.cell { text-align: center; min-width: 88px; font-variant-numeric: tabular-nums; }
  .legend { font-size: 12px; color: var(--muted); margin-top: 8px; }
  .bars { display: grid; gap: 8px; }
  .bar-row { display: grid; grid-template-columns: 280px 1fr 90px; gap: 8px; align-items: center; font-size: 12px; }
  .bar-track { background: #e6eef5; border-radius: 99px; height: 10px; overflow: hidden; }
  .bar-fill { height: 100%; background: linear-gradient(90deg, var(--navy), var(--teal)); }
  .two { display: grid; grid-template-columns: 1fr 1fr; gap: 16px; }
  @media (max-width: 980px) { .two { grid-template-columns: 1fr; } .bar-row { grid-template-columns: 1fr; } }
  .pill { display: inline-block; padding: 2px 8px; border-radius: 99px; font-size: 12px; }
  .up { background: #e5f6ec; color: var(--good); }
  .down { background: #fde8e8; color: var(--bad); }
  .flat { background: #eef2f6; color: var(--muted); }
  .note { font-size: 13px; color: var(--muted); line-height: 1.45; }
</style>
</head>
<body>
<header>
  <h1>Origin–Destination pair-wise yield</h1>
  <p>Month-wise analysis of base yield (basic freight ÷ charged weight) and total yield (total freight ÷ charged weight). Pair yields are volume-weighted: sum of freight divided by sum of weight.</p>
</header>
<div class="wrap">
  <section class="kpis" id="kpis"></section>
  <section class="panel">
    <h2>Filters</h2>
    <div class="filters">
      <label>Month
        <select id="month"></select>
      </label>
      <label>Origin zone
        <select id="origin"></select>
      </label>
      <label>Destination zone
        <select id="dest"></select>
      </label>
      <label>Yield metric
        <select id="metric">
          <option value="base_yield">Base yield</option>
          <option value="total_yield" selected>Total yield</option>
          <option value="vas_yield">VAS yield</option>
        </select>
      </label>
      <button type="button" id="reset">Reset</button>
    </div>
    <p class="note" style="margin:12px 0 0">Heatmap and table follow the filters. Network KPIs always show the filtered subset. Sparse pairs (few shipments) can swing yield; check shipment count before acting on a number.</p>
  </section>
  <section class="panel">
    <h2 id="heatTitle">Yield heatmap</h2>
    <div class="heat" id="heatmap"></div>
    <div class="legend" id="heatLegend"></div>
  </section>
  <section class="two">
    <div class="panel">
      <h2>Network yield by month</h2>
      <div class="bars" id="monthBars"></div>
    </div>
    <div class="panel">
      <h2>Largest OD pairs by weight</h2>
      <div class="bars" id="pairBars"></div>
    </div>
  </section>
  <section class="panel">
    <h2>OD × month yield table</h2>
    <div class="heat" style="max-height: 540px; overflow:auto">
      <table id="grid">
        <thead></thead>
        <tbody></tbody>
      </table>
    </div>
  </section>
</div>
<script>
const DATA = __DATA__;
const fmt = (n, d=2) => n == null || Number.isNaN(n) ? "—" : Number(n).toLocaleString(undefined, {maximumFractionDigits:d, minimumFractionDigits:d});
const fmtInt = (n) => n == null ? "—" : Number(n).toLocaleString(undefined, {maximumFractionDigits:0});
const $ = (id) => document.getElementById(id);

function fillSelect(id, values, allLabel) {
  const el = $(id);
  el.innerHTML = "";
  const all = document.createElement("option");
  all.value = ""; all.textContent = allLabel;
  el.appendChild(all);
  values.forEach(v => {
    const o = document.createElement("option");
    o.value = v; o.textContent = v; el.appendChild(o);
  });
}
fillSelect("month", DATA.months, "All months");
fillSelect("origin", DATA.zones, "All origins");
fillSelect("dest", DATA.zones, "All destinations");

function filteredRows() {
  const m = $("month").value, o = $("origin").value, d = $("dest").value;
  return DATA.od_month.filter(r =>
    (!m || r.month === m) && (!o || r.origin_zone === o) && (!d || r.destination_zone === d)
  );
}

function color(val, min, max) {
  if (val == null) return "#f3f5f8";
  const t = max === min ? 0.5 : (val - min) / (max - min);
  const r = Math.round(244 + (27 - 244) * t);
  const g = Math.round(199 + (122 - 199) * t);
  const b = Math.round(195 + (77 - 195) * t);
  return `rgb(${r},${g},${b})`;
}

function renderKpis(rows) {
  const w = rows.reduce((s,r)=>s+r.charged_weight,0);
  const b = rows.reduce((s,r)=>s+r.basic_freight,0);
  const t = rows.reduce((s,r)=>s+r.total_freight,0);
  const n = rows.reduce((s,r)=>s+r.shipments,0);
  const cards = [
    ["Shipments", fmtInt(n), "in filtered view"],
    ["Charged weight", fmt(w,0), "sum of shipment weight"],
    ["Base yield", w ? fmt(b/w, 4) : "—", "Σ basic freight / Σ weight"],
    ["Total yield", w ? fmt(t/w, 4) : "—", "Σ total freight / Σ weight"],
    ["VAS yield", w ? fmt((t-b)/w, 4) : "—", "(total − basic) / weight"],
    ["OD × month rows", fmtInt(rows.length), "pairs in the current filter"],
  ];
  $("kpis").innerHTML = cards.map(([l,v,s]) =>
    `<div class="kpi"><div class="lbl">${l}</div><div class="val">${v}</div><div class="sub">${s}</div></div>`
  ).join("");
}

function renderHeat(rows) {
  const metric = $("metric").value;
  const months = $("month").value ? [$("month").value] : DATA.months;
  const origins = [...new Set(rows.map(r => r.origin_zone))];
  const dests = [...new Set(rows.map(r => r.destination_zone))];
  const orderedO = DATA.zones.filter(z => origins.includes(z));
  const orderedD = DATA.zones.filter(z => dests.includes(z));
  const map = {};
  rows.forEach(r => { map[r.origin_zone + "|" + r.destination_zone + "|" + r.month] = r; });
  const vals = rows.map(r => r[metric]).filter(v => v != null);
  const min = Math.min(...vals), max = Math.max(...vals);
  $("heatTitle").textContent = ($("metric").selectedOptions[0].text) + " heatmap (origin × destination" + ($("month").value ? ", " + $("month").value : ", all months combined") + ")";

  // If all months, combine by OD
  const combo = {};
  rows.forEach(r => {
    const k = r.origin_zone + "|" + r.destination_zone;
    if (!combo[k]) combo[k] = {w:0,b:0,t:0,v:0,n:0};
    combo[k].w += r.charged_weight; combo[k].b += r.basic_freight;
    combo[k].t += r.total_freight; combo[k].v += r.vas_freight; combo[k].n += r.shipments;
  });
  function cellMetric(c) {
    if (!c || !c.w) return null;
    if (metric === "base_yield") return c.b / c.w;
    if (metric === "total_yield") return c.t / c.w;
    return c.v / c.w;
  }
  const comboVals = Object.values(combo).map(cellMetric).filter(v => v != null);
  const cmin = Math.min(...comboVals), cmax = Math.max(...comboVals);

  let html = "<table><thead><tr><th class='l'>Origin \\ Dest</th>";
  orderedD.forEach(d => html += `<th>${d.replace(" ZONE RO","").replace(" ZONE","")}</th>`);
  html += "</tr></thead><tbody>";
  orderedO.forEach(o => {
    html += `<tr><th class='l'>${o.replace(" ZONE RO","").replace(" ZONE","")}</th>`;
    orderedD.forEach(d => {
      const c = combo[o + "|" + d];
      const val = cellMetric(c);
      const bg = color(val, cmin, cmax);
      const title = c ? `${o} → ${d}\\nShipments: ${c.n}\\nWeight: ${fmt(c.w,0)}` : "No volume";
      html += `<td class="cell" style="background:${bg}" title="${title}">${val==null?"—":fmt(val,3)}<div class="legend">${c?fmtInt(c.n)+" shpts":""}</div></td>`;
    });
    html += "</tr>";
  });
  html += "</tbody></table>";
  $("heatmap").innerHTML = html;
  $("heatLegend").textContent = `Color scale: low ${fmt(cmin,3)} → high ${fmt(cmax,3)} for ${$("metric").selectedOptions[0].text}.`;
}

function renderMonthBars() {
  const m = DATA.months_summary;
  const max = Math.max(...m.map(r => r.total_yield));
  $("monthBars").innerHTML = m.map(r => `
    <div class="bar-row">
      <div><strong>${r.month}</strong><div class="legend">Base ${fmt(r.base_yield,3)} · Total ${fmt(r.total_yield,3)}</div></div>
      <div class="bar-track"><div class="bar-fill" style="width:${100*r.total_yield/max}%"></div></div>
      <div>${fmt(r.total_yield,3)}</div>
    </div>`).join("");
}

function renderPairBars(rows) {
  const combo = {};
  rows.forEach(r => {
    const k = r.od_pair;
    if (!combo[k]) combo[k] = {w:0,b:0,t:0};
    combo[k].w += r.charged_weight; combo[k].b += r.basic_freight; combo[k].t += r.total_freight;
  });
  const list = Object.entries(combo).map(([k,v]) => ({
    pair: k, w: v.w, base: v.b/v.w, total: v.t/v.w
  })).sort((a,b)=>b.w-a.w).slice(0,10);
  const max = Math.max(...list.map(x => x.w), 1);
  $("pairBars").innerHTML = list.map(x => `
    <div class="bar-row">
      <div title="${x.pair}"><strong>${x.pair.replaceAll(" ZONE RO","").replaceAll(" ZONE","")}</strong>
        <div class="legend">Base ${fmt(x.base,3)} · Total ${fmt(x.total,3)}</div></div>
      <div class="bar-track"><div class="bar-fill" style="width:${100*x.w/max}%"></div></div>
      <div>${fmt(x.w,0)}</div>
    </div>`).join("");
}

let sortKey = "charged_weight";
let sortDir = -1;
function momPill(v) {
  if (v == null) return `<span class="pill flat">—</span>`;
  const cls = v > 0.05 ? "up" : (v < -0.05 ? "down" : "flat");
  const sign = v > 0 ? "+" : "";
  return `<span class="pill ${cls}">${sign}${fmt(v,1)}%</span>`;
}

function renderTable(rows) {
  const cols = [
    ["od_pair","OD pair", true],
    ["month","Month", true],
    ["shipments","Shipments", false],
    ["charged_weight","Weight", false],
    ["basic_freight","Basic freight", false],
    ["total_freight","Total freight", false],
    ["base_yield","Base yield", false],
    ["total_yield","Total yield", false],
    ["vas_yield","VAS yield", false],
    ["total_yield_mom_pct","Total yield MoM", false],
  ];
  const sorted = [...rows].sort((a,b) => {
    const av = a[sortKey], bv = b[sortKey];
    if (av == null && bv == null) return 0;
    if (av == null) return 1;
    if (bv == null) return -1;
    if (typeof av === "string") return sortDir * av.localeCompare(bv);
    return sortDir * (av - bv);
  });
  const thead = `<tr>${cols.map(c => `<th class="${c[2]?"l":""}" data-k="${c[0]}">${c[1]}</th>`).join("")}</tr>`;
  $("grid").querySelector("thead").innerHTML = thead;
  $("grid").querySelector("tbody").innerHTML = sorted.map(r => `<tr>
    <td class="l">${r.od_pair}</td>
    <td class="l">${r.month}</td>
    <td>${fmtInt(r.shipments)}</td>
    <td>${fmt(r.charged_weight,1)}</td>
    <td>${fmt(r.basic_freight,0)}</td>
    <td>${fmt(r.total_freight,0)}</td>
    <td>${fmt(r.base_yield,4)}</td>
    <td>${fmt(r.total_yield,4)}</td>
    <td>${fmt(r.vas_yield,4)}</td>
    <td>${momPill(r.total_yield_mom_pct)}</td>
  </tr>`).join("");
  $("grid").querySelectorAll("th").forEach(th => {
    th.onclick = () => {
      const k = th.dataset.k;
      if (sortKey === k) sortDir *= -1; else { sortKey = k; sortDir = k.includes("yield") || k==="od_pair" || k==="month" ? 1 : -1; }
      renderTable(filteredRows());
    };
  });
}

function render() {
  const rows = filteredRows();
  renderKpis(rows);
  renderHeat(rows);
  renderMonthBars();
  renderPairBars(rows);
  renderTable(rows);
}
["month","origin","dest","metric"].forEach(id => $(id).addEventListener("change", render));
$("reset").onclick = () => { ["month","origin","dest"].forEach(id => $(id).value = ""); $("metric").value = "total_yield"; render(); };
render();
</script>
</body>
</html>
"""
    html = html.replace("__DATA__", json.dumps(payload, allow_nan=False))
    HTML_PATH.write_text(html, encoding="utf-8")


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    shipments = load_shipments(DATA_PATH)
    total_w = float(shipments["charged_weight"].sum())
    total_b = float(shipments["basic_freight"].sum())
    total_t = float(shipments["total_freight"].sum())

    od_month = summarize(shipments, ["origin_zone", "destination_zone", "month"])
    od_month = add_shares(od_month, total_w, total_b, total_t)
    od_month = add_mom(od_month)

    od_all = summarize(shipments, ["origin_zone", "destination_zone"])
    od_all = add_shares(od_all, total_w, total_b, total_t)

    month_sum = summarize(shipments, ["month"])
    month_sum = add_shares(month_sum, total_w, total_b, total_t)
    month_sum = month_sum.sort_values("month")

    origin_month = summarize(shipments, ["origin_zone", "month"])
    origin_month = add_shares(origin_month, total_w, total_b, total_t)
    dest_month = summarize(shipments, ["destination_zone", "month"])
    dest_month = add_shares(dest_month, total_w, total_b, total_t)

    quality = shipments[
        (shipments["basic_freight"] <= 0)
        | (shipments["total_freight"] <= 0)
        | (shipments["total_freight"] < shipments["basic_freight"])
    ].copy()
    quality["issue"] = "zero or inverted freight"
    quality.loc[quality["basic_freight"] <= 0, "issue"] = "zero/negative basic freight"
    quality.loc[quality["total_freight"] <= 0, "issue"] = "zero/negative total freight"
    quality.loc[quality["total_freight"] < quality["basic_freight"], "issue"] = "total freight < basic freight"

    od_month.to_csv(CSV_PATH, index=False)
    build_excel(shipments, od_month, od_all, month_sum, origin_month, dest_month, quality)
    build_html(od_month, od_all, month_sum, shipments)

    print(f"Wrote {EXCEL_PATH}")
    print(f"Wrote {HTML_PATH}")
    print(f"Wrote {CSV_PATH}")
    print(
        f"Network base yield={total_b/total_w:.4f}  total yield={total_t/total_w:.4f}  "
        f"rows={len(od_month)} OD pairs={len(od_all)}"
    )


if __name__ == "__main__":
    main()
