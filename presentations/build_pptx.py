#!/usr/bin/env python3
"""Build Supply Chains Under Uncertainty PowerPoint deck."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import nsmap
from pptx.util import Emu, Inches, Pt

# Colors
INK = RGBColor(0x0C, 0x1A, 0x24)
INK_SOFT = RGBColor(0x1A, 0x2F, 0x3D)
SEA = RGBColor(0x0E, 0x4D, 0x5C)
SEA_BRIGHT = RGBColor(0x1A, 0x7A, 0x8C)
SIGNAL = RGBColor(0xC4, 0x5C, 0x26)
SIGNAL_SOFT = RGBColor(0xE8, 0xA8, 0x7C)
PAPER = RGBColor(0xF2, 0xEF, 0xE6)
WHITE = RGBColor(0xFA, 0xF8, 0xF3)
MUTED = RGBColor(0xA8, 0xB4, 0xBC)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def set_run(run, size, bold=False, color=WHITE, font_name="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name


def add_bg(slide, color=INK):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    # Send to back
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)
    return shape


def add_accent_bar(slide, left, top, width, height, color=SIGNAL):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_textbox(slide, left, top, width, height, text, size=18, bold=False,
                color=WHITE, align=PP_ALIGN.LEFT, font_name="Calibri"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size, bold=bold, color=color, font_name=font_name)
    return box


def add_multiline(slide, left, top, width, height, lines, size=16,
                  color=WHITE, bold=False, space_after=6):
    """lines: list of str or (str, dict) with optional bold/color/size."""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        if isinstance(line, tuple):
            text, opts = line
        else:
            text, opts = line, {}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opts.get("align", PP_ALIGN.LEFT)
        p.space_after = Pt(opts.get("space_after", space_after))
        run = p.add_run()
        run.text = text
        set_run(
            run,
            opts.get("size", size),
            bold=opts.get("bold", bold),
            color=opts.get("color", color),
            font_name=opts.get("font_name", "Calibri"),
        )
    return box


def add_card(slide, left, top, width, height, title, body, accent=SEA_BRIGHT):
    card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = INK_SOFT
    card.line.fill.background()
    add_accent_bar(slide, left, top, Inches(0.08), height, accent)
    add_textbox(
        slide, left + Inches(0.25), top + Inches(0.2),
        width - Inches(0.4), Inches(0.4),
        title, size=16, bold=True, color=WHITE
    )
    add_textbox(
        slide, left + Inches(0.25), top + Inches(0.65),
        width - Inches(0.4), height - Inches(0.85),
        body, size=13, color=MUTED
    )


def add_footer(slide, text):
    add_textbox(
        slide, Inches(0.7), Inches(7.05), Inches(12), Inches(0.3),
        text, size=10, color=MUTED
    )


def add_kicker(slide, text):
    add_textbox(
        slide, Inches(0.7), Inches(0.4), Inches(12), Inches(0.35),
        text.upper(), size=12, bold=True, color=SIGNAL_SOFT
    )


def add_title(slide, text, top=Inches(0.8)):
    add_textbox(
        slide, Inches(0.7), top, Inches(12), Inches(0.9),
        text, size=32, bold=True, color=WHITE, font_name="Calibri"
    )


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # ---- SLIDE 1: Title ----
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_accent_bar(s, 0, 0, Inches(0.15), SLIDE_H, SIGNAL)
    add_textbox(
        s, Inches(0.9), Inches(1.4), Inches(11), Inches(0.4),
        "GLOBAL TRADE BRIEFING  ·  OCTOBER 2026",
        size=14, bold=True, color=SIGNAL_SOFT
    )
    add_textbox(
        s, Inches(0.9), Inches(2.0), Inches(11), Inches(1.6),
        "Supply Chains Under Uncertainty",
        size=44, bold=True, color=WHITE
    )
    add_textbox(
        s, Inches(0.9), Inches(3.8), Inches(10), Inches(1.2),
        "How tariffs, conflict, and chokepoint disruption are rewriting "
        "the rules of global logistics — and what leaders are doing about it.",
        size=18, color=MUTED
    )
    # Meta row
    metas = [
        ("Focus", "Current state & strategic trends"),
        ("Context", "Trade policy · Geopolitics · Physical risk"),
        ("Audience", "Operators, planners & executives"),
    ]
    x = Inches(0.9)
    for label, value in metas:
        add_textbox(s, x, Inches(5.6), Inches(3.6), Inches(0.3),
                    label, size=12, bold=True, color=WHITE)
        add_textbox(s, x, Inches(5.95), Inches(3.6), Inches(0.5),
                    value, size=13, color=MUTED)
        x += Inches(3.9)

    # ---- SLIDE 2: Landscape ----
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_kicker(s, "01 — The landscape")
    add_title(s, "From fixed networks to constant motion")
    add_textbox(
        s, Inches(0.7), Inches(1.75), Inches(11.5), Inches(0.8),
        "Efficiency-first globalization is giving way to adaptive, higher-cost "
        "networks. Trade policy, geopolitical risk, and technology are now "
        "embedded in the cost of doing business.",
        size=16, color=MUTED
    )
    cards = [
        ("Resilience over lowest cost",
         "Boards now treat continuity and optionality as competitive assets — "
         "not temporary insurance after a shock.", SEA_BRIGHT),
        ("Policy as standing cost",
         "Tariffs and trade barriers are no longer a short-lived storm to wait "
         "out; they are structural inputs to sourcing math.", SIGNAL),
        ("Fragmented geography",
         "Multi-country footprints — Mexico, Vietnam, India, and beyond — "
         "replace single-source dependence on China.", SEA_BRIGHT),
    ]
    for i, (title, body, accent) in enumerate(cards):
        add_card(
            s, Inches(0.7) + i * Inches(4.05), Inches(2.9),
            Inches(3.85), Inches(3.2), title, body, accent
        )
    add_footer(s, "Sources: KPMG Biannual Supply Chain Report (Mar 2026); DP World Global Trade Observatory 2026")

    # ---- SLIDE 3: Tariffs ----
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_kicker(s, "02 — Tariffs & trade wars")
    add_title(s, "Tariffs rewired sourcing — they did not end imports")
    stats = [
        ("−28%", "Drop in US imports from China in 2025 after renewed tariff escalation"),
        ("+9%", "Rise in US imports from the rest of the world in the same period"),
        ("~9%", "China’s share of US imports by end-2025 (down from ~22% pre-2018)"),
    ]
    for i, (num, label) in enumerate(stats):
        left = Inches(0.7) + i * Inches(4.05)
        add_accent_bar(s, left, Inches(1.9), Inches(3.7), Inches(0.06), SIGNAL)
        add_textbox(s, left, Inches(2.15), Inches(3.7), Inches(0.8),
                    num, size=40, bold=True, color=WHITE)
        add_textbox(s, left, Inches(3.0), Inches(3.7), Inches(1.0),
                    label, size=13, color=MUTED)
    bullets = [
        "Electronics shifted fast to India and Vietnam; apparel often did not — tariff differentials were too small to force a move.",
        "USMCA carve-outs remain a major mitigator; renegotiation risk could add several points to effective tariff rates.",
        "Most firms still raise prices and delay capital projects rather than fully reshore overnight.",
    ]
    add_multiline(
        s, Inches(0.7), Inches(4.3), Inches(12), Inches(2.2),
        [("•  " + b, {"size": 15, "color": WHITE, "space_after": 10}) for b in bullets]
    )
    add_footer(s, "Sources: PIIE Working Paper 26-14; KPMG 2026 Tariff Survey; Kearney 2026 Reshoring Index")

    # ---- SLIDE 4: Geopolitics ----
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_kicker(s, "03 — Wars & geopolitics")
    add_title(s, "Conflict risk is no longer episodic")
    add_accent_bar(s, Inches(0.7), Inches(1.9), Inches(0.08), Inches(1.3), SEA_BRIGHT)
    add_textbox(
        s, Inches(1.0), Inches(1.9), Inches(6.2), Inches(1.3),
        "Geopolitical risk indices in 2025 reached levels not seen since 2022 — "
        "among the highest since the early 2000s wars.",
        size=16, color=PAPER
    )
    bullets = [
        "Russia–Ukraine continues to scramble energy, grain, and industrial metals corridors.",
        "Middle East tensions keep insurance premiums and route choices volatile.",
        "Export controls and dual-use restrictions harden tech and critical-materials flows.",
    ]
    add_multiline(
        s, Inches(0.7), Inches(3.5), Inches(6.5), Inches(2.5),
        [("•  " + b, {"size": 15, "color": WHITE, "space_after": 12}) for b in bullets]
    )
    add_card(
        s, Inches(7.7), Inches(1.9), Inches(4.9), Inches(2.1),
        "Political distance ≠ supply safety",
        "“Friendly” assembly hubs often still depend on Chinese inputs — "
        "friendshoring can relocate risk rather than eliminate it.",
        SIGNAL
    )
    add_card(
        s, Inches(7.7), Inches(4.2), Inches(4.9), Inches(2.1),
        "Scenario planning is table stakes",
        "Leaders model tariff, conflict, and chokepoint scenarios together — "
        "not as separate risk silos.",
        SEA_BRIGHT
    )
    add_footer(s, "Sources: KPMG Mar 2026; Economist Impact Trade in Transition 2025; Capgemini Reindustrialization 2025")

    # ---- SLIDE 5: Physical disruption ----
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_kicker(s, "04 — Physical disruption")
    add_title(s, "Chokepoints still set the tempo of trade")
    cards = [
        ("Red Sea / Suez",
         "Attacks forced widespread Cape of Good Hope detours — adding roughly "
         "7–12 days and large fuel and capacity penalties on Asia–Europe lanes.",
         SIGNAL),
        ("Panama Canal",
         "Drought and transit limits compressed capacity on Americas–Asia routes, "
         "amplifying backlog risk when other corridors strain.",
         SEA_BRIGHT),
        ("Freight & capacity",
         "Rerouting cuts effective global capacity and lifts rates; disruption "
         "compounds tariff-driven demand shifts.",
         SEA_BRIGHT),
    ]
    for i, (title, body, accent) in enumerate(cards):
        add_card(
            s, Inches(0.7) + i * Inches(4.05), Inches(1.9),
            Inches(3.85), Inches(2.8), title, body, accent
        )
    bullets = [
        "Contingency routing and multi-carrier contracts moved from “nice to have” to operating standard.",
        "Inventory buffers rose again: executives cite stock builds alongside diversification as top 2026 moves.",
    ]
    add_multiline(
        s, Inches(0.7), Inches(5.1), Inches(12), Inches(1.4),
        [("•  " + b, {"size": 15, "color": WHITE, "space_after": 10}) for b in bullets]
    )
    add_footer(s, "Sources: Container News; industry disruption briefings 2025–26; DP World GTO 2026")

    # ---- SLIDE 6: Strategic responses ----
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_kicker(s, "05 — Strategic responses")
    add_title(s, "Four playbooks, one goal: optionality")
    cards = [
        ("China + 1 / + N",
         "Keep China where it still wins on scale; add Vietnam, India, Mexico, "
         "and others for tariff and continuity hedges. Cited by ~51% of trade "
         "execs as a top 2026 shift.", SEA_BRIGHT),
        ("Friendshoring",
         "~36% plan deeper ties with politically aligned partners — useful, "
         "but alliance labels shift with policy.", SEA_BRIGHT),
        ("Nearshoring",
         "Mexico remains pivotal for US-bound goods, yet tariff uncertainty and "
         "USMCA risk have cooled confidence for many.", SEA_BRIGHT),
        ("Reshoring",
         "Formal reshoring plans jumped (to ~26% in KPMG survey), but most say "
         "full moves take 1–3 years amid labor, capital, and tier-2 gaps.", SIGNAL),
    ]
    positions = [
        (Inches(0.7), Inches(1.85)),
        (Inches(6.85), Inches(1.85)),
        (Inches(0.7), Inches(4.35)),
        (Inches(6.85), Inches(4.35)),
    ]
    for (title, body, accent), (left, top) in zip(cards, positions):
        add_card(s, left, top, Inches(5.75), Inches(2.25), title, body, accent)
    add_footer(s, "Sources: DP World Global Trade Observatory 2026; Kearney Reshoring Index 2026; KPMG Tariff Survey 2026")

    # ---- SLIDE 7: Data ----
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_kicker(s, "06 — What the data says")
    add_title(s, "Confidence in trade — with eyes wide open")
    stats = [
        ("94%", "of trade executives expect 2026 growth to match or beat 2025"),
        ("90%", "expect trade barriers to rise or stay elevated"),
        ("68%", "of firms delayed or postponed investments under tariff pressure"),
    ]
    for i, (num, label) in enumerate(stats):
        left = Inches(0.7) + i * Inches(4.05)
        add_accent_bar(s, left, Inches(1.9), Inches(3.7), Inches(0.06), SIGNAL)
        add_textbox(s, left, Inches(2.15), Inches(3.7), Inches(0.8),
                    num, size=40, bold=True, color=WHITE)
        add_textbox(s, left, Inches(3.0), Inches(3.7), Inches(1.0),
                    label, size=13, color=MUTED)
    bullets = [
        "US manufacturing imports hit a four-year high despite record reshoring talk — tariffs redirected origins more than they eliminated import dependence.",
        "Margin pressure is structural: price pass-through and selective inventory remain the common near-term tools.",
        "Digital visibility, supplier mapping, and FTA-aware routing are rising as growth enablers, not only as risk tools.",
    ]
    add_multiline(
        s, Inches(0.7), Inches(4.3), Inches(12), Inches(2.2),
        [("•  " + b, {"size": 15, "color": WHITE, "space_after": 10}) for b in bullets]
    )
    add_footer(s, "Sources: DP World GTO 2026; Kearney 2026; KPMG 2026 Tariff Survey")

    # ---- SLIDE 8: Implications ----
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_kicker(s, "07 — Implications for leaders")
    add_title(s, "Design for disruption, not for calm")
    bullets = [
        "Treat tariffs as a permanent planning variable — model rate paths, carve-outs, and retaliation in every sourcing case.",
        "Map multi-tier exposure — assembly outside China still fails if critical inputs and tooling do not.",
        "Invest in dual routes and dual suppliers — chokepoint and conflict risk will keep capacity lumpy.",
        "Prefer flexible capital — smaller, staged bets beat large irreversible bets while policy stays fluid.",
        "Make resilience a P&L story — boards reward continuity that protects revenue, not only cost wins.",
    ]
    add_multiline(
        s, Inches(0.7), Inches(1.9), Inches(7.2), Inches(3.8),
        [("•  " + b, {"size": 15, "color": WHITE, "space_after": 14}) for b in bullets]
    )
    # Closing quote panel
    panel = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(8.2), Inches(1.9), Inches(4.4), Inches(3.8)
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = INK_SOFT
    panel.line.fill.background()
    add_accent_bar(s, Inches(8.2), Inches(1.9), Inches(4.4), Inches(0.08), SIGNAL)
    add_textbox(
        s, Inches(8.5), Inches(2.3), Inches(3.9), Inches(2.2),
        "The winners of this cycle will not be the cheapest networks — "
        "they will be the ones that can reconfigure fastest when the next shock lands.",
        size=16, color=PAPER
    )
    add_textbox(
        s, Inches(8.5), Inches(4.8), Inches(3.9), Inches(0.5),
        "Resilience is now a growth strategy.",
        size=15, bold=True, color=SIGNAL_SOFT
    )
    add_footer(
        s,
        "Briefing synthesized from KPMG, PIIE, Kearney, DP World, Capgemini, "
        "and Economist Impact research, 2025–2026."
    )

    out = Path(__file__).resolve().parent / "Supply_Chains_Under_Uncertainty.pptx"
    prs.save(out)
    print(f"Wrote {out}")
    return out


if __name__ == "__main__":
    build()
